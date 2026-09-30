"""Exercise real shared validation through BFQ execution and operator boundaries."""

import hashlib
import json
import logging
import sys

from pathlib import Path
from unittest.mock import Mock

import flowcell_manager.flowcell_manager as manager
import pytest

from bcl2fastq_pipeline.state import (
    ExecutionLeaseError,
    FlowcellStateStore,
    cleanup_plan,
    new_state,
)
from openpyxl import load_workbook
from test_state_integration import (
    RUN_ID,
    completed_state,
    configured_bfq,
    write_fastq,
    write_inputs,
)

from bcl2fastq_pipeline import (
    afterFastq,
    cli,
    findFlowCells,
    makeFastq,
    misc,
    notifications,
    preflight,
)


def change_customer(directory, rows, headers=None):
    path = directory / "Sample-Submission-Form.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["Sample-Submission-Form"]
    sheet.delete_rows(15, sheet.max_row)
    for index, header in enumerate(headers or ["Unique Sample ID", "Sample Group"], start=1):
        sheet.cell(15, index, header)
    for row_number, values in enumerate(rows, start=16):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
    workbook.save(path)


def change_lab(directory, rows):
    path = directory / "Sample-Submission-Form.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["INFO (GCF-lab only)"]
    sheet.delete_rows(2, sheet.max_row)
    for values in rows:
        sheet.append(values)
    workbook.save(path)


def replace_data(directory, data):
    path = directory / "SampleSheet.csv"
    custom = path.read_text().split("[Data]", 1)[0]
    path.write_text(custom + "[Data]\n" + data)


def snapshot(directory):
    return {
        str(path.relative_to(directory)): (path.read_bytes(), path.stat().st_mtime_ns)
        for path in directory.rglob("*")
        if path.is_file()
    }


def simulated_execution(cfg, monkeypatch):
    """Replace compute/delivery only: selection, validator and state are real."""
    calls = []
    monkeypatch.setattr(misc, "enoughFreeSpace", lambda: True)
    monkeypatch.setattr(cli, "_prepare_workflow", lambda *_args: calls.append("workflow"))
    monkeypatch.setattr(makeFastq, "bcl2fq", lambda: (calls.append("demux") or ("test", "1")))
    monkeypatch.setattr(makeFastq, "rename_fastqs", lambda: calls.append("rename"))
    monkeypatch.setattr(
        afterFastq, "md5sum_worker", lambda *_args, **_kwargs: calls.append("checksum")
    )
    monkeypatch.setattr(afterFastq, "analysis_steps", lambda: calls.append("analysis"))
    monkeypatch.setattr(cli, "_run_reporting", lambda *_args: ["GCF-2026-001"])
    monkeypatch.setattr(afterFastq, "finalize", lambda: calls.append("finalize"))
    monkeypatch.setattr(findFlowCells, "markFinished", lambda: ["GCF-2026-001"])
    monkeypatch.setattr(notifications, "make_payload", lambda *_args, **_kwargs: {"run_id": RUN_ID})
    monkeypatch.setattr(cli.sequencing_delivery, "ensure_report", lambda *_a, **_kw: False)
    monkeypatch.setattr(cli.notification_delivery, "deliver_pending", lambda *_a, **_kw: True)
    monkeypatch.setattr(misc, "send_error_report", Mock())
    return calls


@pytest.mark.parametrize("boundary", ["demultiplexing", "analysis"])
@pytest.mark.parametrize(
    "problem", ["missing-id", "malformed-workbook", "bad-header", "merge-loss", "ambiguous-mapping"]
)
def test_invalid_pair_blocks_execution_and_preserves_structured_context(
    tmp_path, monkeypatch, boundary, problem
):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source if boundary == "demultiplexing" else output)
    directory = source if boundary == "demultiplexing" else output
    if problem == "missing-id":
        replace_data(directory, "Sample_ID,Sample_Project,index\nmissing,GCF-2026-001,ACGT\n")
    elif problem == "malformed-workbook":
        (directory / "Sample-Submission-Form.xlsx").write_bytes(b"not a workbook")
    elif problem == "bad-header":
        change_customer(directory, [["sample", "group"]], headers=["Wrong header", "Sample Group"])
    elif problem == "merge-loss":
        change_customer(directory, [["sample", "group"], ["extra", "group"]])
        change_lab(directory, [["extra"]])
    else:
        change_customer(directory, [["sample", "first"], ["sample", "second"]])
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(
        new_state(
            RUN_ID,
            source,
            output,
            origin="new" if boundary == "demultiplexing" else "restored_legacy_fastq",
            start_stage=boundary,
            cfg=cfg,
        )
    )
    calls = simulated_execution(cfg, monkeypatch)
    real_validate = preflight.validate_selection

    def validate_under_lease(selection):
        with pytest.raises(ExecutionLeaseError), store.execution_lease(RUN_ID):
            pass
        return real_validate(selection)

    monkeypatch.setattr(preflight, "validate_selection", validate_under_lease)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)

    state = store.read(RUN_ID)
    assert calls == []
    assert state["status"] == "failed"
    assert state["current_stage"] == boundary
    assert "Input preflight failed" in state["last_error"]["summary"]
    report = state["stages"][boundary]["metadata"]["input_preflight"]
    assert report["status"] == "failed"
    assert report["validator"]["version"]
    assert report["errors"]
    assert report == json.loads(
        (cfg.static.paths.report_dir / f"{RUN_ID}.input-preflight.json").read_text()
    )
    assert report == state["attempts"][-1]["input_preflight"][boundary]
    assert {item["sha256"] for item in report["inputs"]} == {
        hashlib.sha256((output / name).read_bytes()).hexdigest()
        for name in ("SampleSheet.csv", "Sample-Submission-Form.xlsx")
    }


def test_valid_subset_extra_metadata_and_multilane_reaches_demux(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    change_customer(source, [["sample", "group"], ["extra", "extra-group"]])
    change_lab(source, [["sample"], ["extra"]])
    replace_data(
        source,
        "Lane,Sample_ID,Sample_Project,index\n1,sample,GCF-2026-001,ACGT\n2,sample,GCF-2026-001,ACGT\n",
    )
    assert findFlowCells.flowCellProcessed() is False
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    calls = simulated_execution(cfg, monkeypatch)

    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)

    state = store.read(RUN_ID)
    assert state["status"] == "completed"
    assert "demux" in calls
    summary = state["stages"]["demultiplexing"]["metadata"]["input_preflight"]["summary"]
    assert summary["planned_sample_count"] == 1
    assert summary["extra_submission_sample_count"] == 1
    assert summary["samplesheet_rows"] == 2


def test_manual_validation_before_discovery_is_read_only(tmp_path, capsys):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    cfg.static.paths.manager_dir.rmdir()
    before = snapshot(tmp_path)
    result = manager.validate_flowcell(flowcell=RUN_ID)
    assert result.ok
    assert snapshot(tmp_path) == before
    assert not output.exists()
    assert not cfg.static.paths.manager_dir.exists()
    assert str(source / "SampleSheet.csv") in capsys.readouterr().out


def test_manual_validation_does_not_recover_running_state(tmp_path):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(output)
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(new_state(RUN_ID, source, output, origin="new", start_stage="demultiplexing"))
    store.begin_attempt(RUN_ID)
    before = snapshot(tmp_path)
    manager.validate_flowcell(flowcell=RUN_ID)
    assert snapshot(tmp_path) == before
    assert store.read(RUN_ID)["status"] == "running"


def test_manual_invalid_command_returns_nonzero_without_creating_state(
    tmp_path, monkeypatch, capsys
):
    cfg, source, _output = configured_bfq(tmp_path)
    write_inputs(source)
    replace_data(source, "Sample_ID,Sample_Project\nmissing,GCF-2026-001\n")
    cfg.static.paths.manager_dir.rmdir()
    before = snapshot(tmp_path)
    monkeypatch.setattr(sys, "argv", ["flowcell-manager", "validate", RUN_ID])
    with pytest.raises(SystemExit) as error:
        manager.main()
    assert error.value.code == 1
    assert "missing" in capsys.readouterr().err
    assert snapshot(tmp_path) == before
    assert not cfg.static.paths.manager_dir.exists()


@pytest.mark.parametrize("curated_file", ["SampleSheet.csv", "Sample-Submission-Form.xlsx"])
def test_partial_curated_pair_is_preserved_and_manual_agrees_with_execution(
    tmp_path, monkeypatch, curated_file
):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source, "-instrument")
    write_inputs(output, "-curated")
    missing = (
        "Sample-Submission-Form.xlsx" if curated_file == "SampleSheet.csv" else "SampleSheet.csv"
    )
    (output / missing).unlink()
    curated_bytes = (output / curated_file).read_bytes()
    manual = manager.validate_flowcell(flowcell=RUN_ID).to_dict()
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(new_state(RUN_ID, source, output, origin="new", start_stage="demultiplexing"))
    simulated_execution(cfg, monkeypatch)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    automatic = store.read(RUN_ID)["stages"]["demultiplexing"]["metadata"]["input_preflight"]
    assert automatic["status"] == manual["status"] == "passed"
    assert automatic["summary"] == manual["summary"]
    assert {item["sha256"] for item in automatic["inputs"]} == {
        item["sha256"] for item in manual["inputs"]
    }
    assert (output / curated_file).read_bytes() == curated_bytes
    assert (output / missing).read_bytes() == (source / missing).read_bytes()


def test_malformed_curated_sheet_is_not_replaced_by_good_instrument_sheet(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    write_inputs(output)
    (output / "SampleSheet.csv").write_bytes(b"\xff invalid UTF8")
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(new_state(RUN_ID, source, output, origin="new", start_stage="demultiplexing"))
    calls = simulated_execution(cfg, monkeypatch)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    assert calls == []
    assert store.read(RUN_ID)["status"] == "failed"
    assert (output / "SampleSheet.csv").read_bytes() == b"\xff invalid UTF8"


def test_source_selection_preserves_custom_options_opt_in(tmp_path):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    (source / "SampleSheet.csv").rename(source / "SampleSheet-bfq.csv")
    (source / "SampleSheet.csv").write_text(
        "[Data]\nSample_ID,Sample_Project\noriginal,GCF-2026-001\n"
    )
    selection = preflight.select_run_inputs(source, output)
    assert selection.sample_sheet.name == "SampleSheet-bfq.csv"
    assert findFlowCells.flowCellProcessed() is False
    findFlowCells.newFlowCell()
    assert cfg.run.custom["Libprep"] == "Illumina DNA Prep"
    assert (output / "SampleSheet.csv").read_bytes() == (
        source / "SampleSheet-bfq.csv"
    ).read_bytes()


def test_non_bfq_instrument_run_is_not_automatically_opted_in(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    (source / "SampleSheet.csv").write_text(
        "[Data]\nSample_ID,Sample_Project\nsample,GCF-2026-001\n"
    )
    assert findFlowCells.flowCellProcessed() is False
    calls = simulated_execution(cfg, monkeypatch)
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    assert calls == []
    assert store.read(RUN_ID)["status"] == "queued"
    assert not output.exists()


@pytest.mark.parametrize("alternate", [False, True])
def test_malformed_opted_in_instrument_sheet_fails_visibly(tmp_path, monkeypatch, alternate):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    sheet = source / ("SampleSheet-bfq.csv" if alternate else "SampleSheet.csv")
    if alternate:
        (source / "SampleSheet.csv").write_text(
            "[Data]\nSample_ID,Sample_Project\nsample,GCF-2026-001\n"
        )
    content = b"[CustomOptions]\nLibprep,Illumina DNA Prep\n[Data]\n\xff invalid UTF8"
    sheet.write_bytes(content)
    assert findFlowCells.flowCellProcessed() is False
    calls = simulated_execution(cfg, monkeypatch)
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    assert calls == []
    assert store.read(RUN_ID)["status"] == "failed"
    assert (output / "SampleSheet.csv").read_bytes() == content


@pytest.mark.parametrize("boundary", ["demultiplexing", "analysis"])
@pytest.mark.parametrize("refresh", [False, True])
def test_invalid_rerun_does_not_remove_existing_products_or_change_state(
    tmp_path, boundary, refresh
):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    write_inputs(output)
    bad_directory = source if refresh else output
    replace_data(bad_directory, "Sample_ID,Sample_Project\nmissing,GCF-2026-001\n")
    (output / "multiqc_GCF-2026-001.html").write_text("existing report")
    (output / "GCF-2026-001.7za").write_bytes(b"existing archive")
    write_fastq(output / "GCF-2026-001/sample_R1.fastq.gz")
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(completed_state(cfg, source, output))
    before = snapshot(output), snapshot(source), store.state_path(RUN_ID).read_bytes()
    with pytest.raises(preflight.PreflightValidationError):
        manager.rerun_flowcell(
            flowcell=RUN_ID, from_stage=boundary, force=True, refresh_inputs=refresh
        )
    assert (snapshot(output), snapshot(source), store.state_path(RUN_ID).read_bytes()) == before


def test_refresh_preview_does_not_copy_and_refresh_rerun_uses_instrument(tmp_path):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source, "-instrument")
    write_inputs(output, "-curated")
    replace_data(output, "Sample_ID,Sample_Project\nmissing,GCF-2026-001\n")
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(completed_state(cfg, source, output))
    before = snapshot(tmp_path)
    assert manager.validate_flowcell(flowcell=RUN_ID, refresh_inputs=True).ok
    assert snapshot(tmp_path) == before
    queued = manager.rerun_flowcell(
        flowcell=RUN_ID, from_stage="analysis", force=True, refresh_inputs=True
    )
    assert queued["status"] == "queued"
    assert (output / "SampleSheet.csv").read_bytes() == (source / "SampleSheet.csv").read_bytes()


def test_corrected_inputs_retry_without_state_edit_and_do_not_reuse_result(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    replace_data(source, "Sample_ID,Sample_Project\nmissing,GCF-2026-001\n")
    assert findFlowCells.flowCellProcessed() is False
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    calls = simulated_execution(cfg, monkeypatch)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    failed = store.read(RUN_ID)["stages"]["demultiplexing"]["metadata"]["input_preflight"]
    assert failed["status"] == "failed"
    replace_data(output, "Sample_ID,Sample_Project\nsample,GCF-2026-001\n")
    manager.rerun_flowcell(flowcell=RUN_ID, from_stage="demultiplexing", force=True)
    cfg.run.begin(source, cfg.static.paths)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    state = store.read(RUN_ID)
    assert state["status"] == "completed"
    assert "demux" in calls
    passed = state["stages"]["demultiplexing"]["metadata"]["input_preflight"]
    assert passed["status"] == "passed"
    assert passed["inputs"] != failed["inputs"]
    assert state["attempts"][0]["input_preflight"]["demultiplexing"] == failed


def test_restored_fastqs_validate_before_analysis(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(output)
    write_fastq(output / "GCF-2026-001/sample_R1.fastq.gz")
    assert findFlowCells.flowCellProcessed() is False
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    assert store.read(RUN_ID)["origin"] == "restored_legacy_fastq"
    replace_data(output, "Sample_ID,Sample_Project\nmissing,GCF-2026-001\n")
    calls = simulated_execution(cfg, monkeypatch)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    assert calls == []
    assert store.read(RUN_ID)["current_stage"] == "analysis"
    assert store.read(RUN_ID)["status"] == "failed"


@pytest.mark.parametrize("boundary", ["demultiplexing", "analysis"])
def test_daemon_revalidates_inputs_changed_after_operator_queued_run(
    tmp_path, monkeypatch, boundary
):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(output)
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(completed_state(cfg, source, output))
    assert (
        manager.rerun_flowcell(flowcell=RUN_ID, from_stage=boundary, force=True)["status"]
        == "queued"
    )
    replace_data(output, "Sample_ID,Sample_Project\nchanged-after-queue,GCF-2026-001\n")
    calls = simulated_execution(cfg, monkeypatch)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    assert calls == []
    assert store.read(RUN_ID)["status"] == "failed"


def test_invalid_initialization_keeps_existing_products_and_creates_no_state(tmp_path):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(output)
    (output / "Sample-Submission-Form.xlsx").write_bytes(b"broken workbook")
    (output / "existing.7za").write_bytes(b"existing result")
    before = snapshot(tmp_path)
    with pytest.raises(preflight.PreflightValidationError):
        manager.initialize_flowcell(flowcell=RUN_ID, from_stage="demultiplexing", force=True)
    assert snapshot(tmp_path) == before
    assert not FlowcellStateStore(cfg.static.paths.manager_dir).exists(RUN_ID)


def test_legacy_filenames_survive_initialization_cleanup_as_canonical_inputs(tmp_path):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(output)
    (output / "SampleSheet.csv").rename(output / "SampleSheet-curated.csv")
    (output / "Sample-Submission-Form.xlsx").rename(output / "old-Sample-Submission-Form.xlsx")
    expected = preflight.select_run_inputs(source, output)
    expected_bytes = expected.sample_sheet.read_bytes(), expected.submission_form.read_bytes()
    assert (
        manager.initialize_flowcell(flowcell=RUN_ID, from_stage="demultiplexing", force=True)[
            "status"
        ]
        == "queued"
    )
    assert (output / "SampleSheet.csv").read_bytes() == expected_bytes[0]
    assert (output / "Sample-Submission-Form.xlsx").read_bytes() == expected_bytes[1]
    assert not expected.sample_sheet.exists()
    assert not expected.submission_form.exists()


def test_fresh_run_missing_submission_form_records_actionable_failure(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    (source / "Sample-Submission-Form.xlsx").unlink()
    assert findFlowCells.flowCellProcessed() is False
    calls = simulated_execution(cfg, monkeypatch)
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    state = store.read(RUN_ID)
    assert state["status"] == "failed"
    assert "Submission-Form.xlsx" in state["last_error"]["summary"]
    assert calls == []


def test_warning_only_pair_does_not_block_execution(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    change_customer(
        source,
        [["sample", "group", "GCF-2026-OTHER"]],
        headers=["Unique Sample ID", "Sample Group", "Project"],
    )
    assert findFlowCells.flowCellProcessed() is False
    calls = simulated_execution(cfg, monkeypatch)
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    state = store.read(RUN_ID)
    report = state["stages"]["demultiplexing"]["metadata"]["input_preflight"]
    assert report["warnings"]
    assert report["status"] == "passed"
    assert state["status"] == "completed"
    assert "demux" in calls


@pytest.mark.parametrize(
    "boundary,removed", [("analysis", True), ("reporting", False), ("finalization", False)]
)
def test_actual_analysis_sample_summary_is_invalidated_only_with_analysis(
    tmp_path, boundary, removed
):
    output = tmp_path / RUN_ID
    output.mkdir()
    summary = output / "configmaker-analysis-GCF-2026-001.json"
    summary.write_text('{"discovered_sample_count": 1}')
    assert (summary in cleanup_plan(output, boundary)) is removed


@pytest.mark.parametrize("valid", [False, True])
def test_sidecar_write_failure_keeps_durable_report_and_original_diagnosis(
    tmp_path, monkeypatch, valid
):
    cfg, source, output = configured_bfq(tmp_path)
    write_inputs(source)
    if not valid:
        replace_data(source, "Sample_ID,Sample_Project\nmissing,GCF-2026-001\n")
    assert findFlowCells.flowCellProcessed() is False
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    calls = simulated_execution(cfg, monkeypatch)
    original_write = Path.write_text

    def fail_sidecar(path, *args, **kwargs):
        if ".input-preflight.json." in path.name:
            raise OSError("simulated sidecar failure")
        return original_write(path, *args, **kwargs)

    monkeypatch.setattr(Path, "write_text", fail_sidecar)
    cli._run_state_backed_flowcell(cfg, store, logging.getLogger("preflight-test"), prepare=True)
    state = store.read(RUN_ID)
    report = state["stages"]["demultiplexing"]["metadata"]["input_preflight"]
    assert report["report_path"] is None
    assert report["report_write_error"] == "simulated sidecar failure"
    assert report["status"] == ("passed" if valid else "failed")
    if valid:
        assert "demux" in calls
        assert state["status"] == "completed"
    else:
        assert calls == []
        assert state["status"] == "failed"
        assert "Input preflight failed" in state["last_error"]["summary"]
        assert "missing" in state["last_error"]["summary"]
