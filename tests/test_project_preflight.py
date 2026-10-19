"""Curated multi-pair selection uses project keys and validates every pair."""

import shutil

import pytest

from openpyxl import load_workbook
from support.fixtures import write_inputs

from bcl2fastq_pipeline import preflight


def pair(directory, project, sample, index):
    temporary = directory / project
    write_inputs(temporary)
    sheet = directory / f"SampleSheet_{project}.csv"
    sheet.write_text(
        "[CustomOptions]\nLibprep,Illumina DNA Prep\nUser,test\n"
        f"[Data]\nSample_ID,Sample_Project,index\n{sample},{project},{index}\n"
    )
    form = directory / f"{project}_Sample-Submission-Form.xlsx"
    workbook = load_workbook(temporary / "Sample-Submission-Form.xlsx")
    workbook["Sample-Submission-Form"].cell(16, 1, sample)
    workbook["INFO (GCF-lab only)"].cell(2, 1, sample)
    workbook.save(form)
    shutil.rmtree(temporary)
    return sheet, form


def test_refresh_uses_selected_instrument_pair_and_curated_unselected_pair(tmp_path):
    source, output = tmp_path / "source", tmp_path / "output"
    source.mkdir()
    output.mkdir()
    for directory in (source, output):
        pair(directory, "GCF-2099-001", "alpha", "AAAACCCC")
        pair(directory, "GCF-2099-002", "beta", "TTTTGGGG")
    selected = preflight.select_run_inputs(source, output, refresh=True, projects=["GCF-2099-002"])
    assert selected.multi_pair
    assert {item.project_id: item.sample_sheet.parent for item in selected.pairs} == {
        "GCF-2099-001": output,
        "GCF-2099-002": source,
    }
    assert preflight.validate_selection(selected).ok


def test_unselected_invalid_pair_blocks_refresh_before_copy(tmp_path):
    source, output = tmp_path / "source", tmp_path / "output"
    source.mkdir()
    output.mkdir()
    for directory in (source, output):
        pair(directory, "GCF-2099-001", "alpha", "AAAACCCC")
        pair(directory, "GCF-2099-002", "beta", "TTTTGGGG")
    broken = output / "GCF-2099-001_Sample-Submission-Form.xlsx"
    broken.write_bytes(b"invalid curated workbook")
    before = {path.name: path.read_bytes() for path in output.iterdir()}
    with pytest.raises(preflight.PreflightValidationError):
        preflight.require_valid_inputs(source, output, refresh=True, projects=["GCF-2099-002"])
    assert {path.name: path.read_bytes() for path in output.iterdir()} == before


def test_copy_retains_project_filenames_and_duplicate_provenance(tmp_path):
    source, output = tmp_path / "source", tmp_path / "output"
    source.mkdir()
    first, _ = pair(source, "GCF-2099-001", "alpha", "AAAACCCC")
    pair(source, "GCF-2099-002", "beta", "TTTTGGGG")
    shutil.copy2(first, source / "SampleSheet.csv")
    selected = preflight.select_run_inputs(source, output)
    copied = preflight.copy_run_inputs(selected, output)
    assert {item.project_id for item in copied.pairs} == {"GCF-2099-001", "GCF-2099-002"}
    assert not (output / "SampleSheet.csv").exists()
    assert "SampleSheet.csv" in (output / "project-inputs.json").read_text()
    assert all(item.sample_sheet.parent == output for item in copied.pairs)


def test_legacy_curated_canonical_pair_remains_authoritative(tmp_path):
    source, output = tmp_path / "source", tmp_path / "output"
    source.mkdir()
    write_inputs(output)
    pair(source, "GCF-2099-001", "alpha", "AAAACCCC")
    pair(source, "GCF-2099-002", "beta", "TTTTGGGG")
    selected = preflight.select_run_inputs(source, output)
    assert not selected.multi_pair
    assert selected.sample_sheet == output / "SampleSheet.csv"
