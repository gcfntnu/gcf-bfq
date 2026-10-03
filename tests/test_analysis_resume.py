"""Resume safety and orchestration without scientific tools or operational data."""

import json
import shutil

from unittest.mock import Mock

import flowcell_manager.flowcell_manager as manager
import pandas as pd
import pytest
import yaml

from bcl2fastq_pipeline.state import ExecutionLeaseError, FlowcellStateStore, StateConflictError
from openpyxl import load_workbook
from support.fixtures import RUN_ID, completed_state, configured_bfq, write_fastq, write_inputs

from bcl2fastq_pipeline import afterFastq, analysis_resume, analysis_snapshots, preflight

PROJECT = "GCF-2026-001"


@pytest.fixture
def prepared(tmp_path, monkeypatch):
    cfg, source, output = configured_bfq(tmp_path)
    monkeypatch.setenv("TMPDIR", str(tmp_path / "scratch"))
    monkeypatch.setenv("SINGULARITY_CACHEDIR", str(tmp_path / "cache"))
    write_inputs(output)
    fastq = output / PROJECT / "sample_R1.fastq.gz"
    write_fastq(fastq)
    state = completed_state(cfg, source, output)
    work = analysis_snapshots.workdir_path(RUN_ID, PROJECT)
    work.mkdir(parents=True)
    identity = analysis_snapshots.identify_workdir(work, RUN_ID, PROJECT)
    state["stages"]["analysis"]["metadata"] = {
        "workflow": "singlecell",
        "workdirs": {PROJECT: identity},
    }
    config = {
        "project_id": [PROJECT],
        "workflow": "singlecell",
        "samples": {
            "sample": {
                "Sample_ID": "sample",
                "Flowcell_Name": RUN_ID,
                "Flowcell_ID": RUN_ID.split("_")[-1],
            }
        },
    }
    (work / "config.yaml").write_text(yaml.safe_dump(config))
    (work / "Snakefile").write_text("# retained local edit\n")
    (work / "pep").mkdir()
    (work / "pep/pep_config.yaml").write_text("sample_table: sample_table.csv\n")
    (work / "pep/sample_table.csv").write_text("sample_name\nsample\n")
    workflow = work / "src/gcf-workflows/singlecell"
    workflow.mkdir(parents=True)
    (workflow / "singlecell.smk").write_text("# locally repaired workflow\n")
    (workflow.parent / "libprep.config").write_text("retained config\n")
    raw = work / "data/raw/fastq"
    raw.mkdir(parents=True)
    (raw / fastq.name).symlink_to(fastq)
    (work / ".snakemake").mkdir()
    (work / ".snakemake/metadata").write_text("provenance\n")
    result = work / "data/tmp/singlecell/bfq"
    result.mkdir(parents=True)
    (result / f"multiqc_{PROJECT}.html").write_text("current report")
    (result / ".multiqc_config.yaml").write_text("{}\n")
    (work / "data/tmp/sample_info.tsv").write_text("sample\tgroup\n")
    store = FlowcellStateStore(cfg.static.paths.manager_dir)
    store.create(state)
    return cfg, store, output, work


def bytes_and_times(root):
    return {
        str(p.relative_to(root)): (p.read_bytes(), p.stat().st_mtime_ns)
        for p in root.rglob("*")
        if p.is_file()
    }


def queue(**kw):
    return manager.rerun_flowcell(
        flowcell=RUN_ID, from_stage="analysis", resume=True, force=True, **kw
    )


def test_preview_queue_and_attempt_preserve_workspace_and_bind_context(prepared, capsys):
    cfg, store, output, work = prepared
    (output / "old.7za").write_bytes(b"existing delivery")
    before = bytes_and_times(work), bytes_and_times(output), store.state_path(RUN_ID).read_bytes()
    queue(dry_run=True)
    assert "Manager preview only" in capsys.readouterr().out
    assert (
        bytes_and_times(work),
        bytes_and_times(output),
        store.state_path(RUN_ID).read_bytes(),
    ) == before
    queued = queue(reason="local repair")
    assert queued["status"] == "queued"
    assert queued["stages"]["demultiplexing"]["status"] == "completed"
    assert queued["stages"]["analysis"]["metadata"] == {}
    assert queued["stages"]["reporting"]["status"] == "pending"
    assert (bytes_and_times(work), bytes_and_times(output)) == before[:2]
    restarted = FlowcellStateStore(cfg.static.paths.manager_dir)
    attempt = restarted.begin_attempt(RUN_ID)
    context = attempt["attempts"][-1]["analysis_resume"]
    assert context == queued["restart_request"]["analysis_resume"]
    analysis_resume.verify(context, analysis_resume.inspect(attempt))
    (work / "config.yaml").write_text((work / "config.yaml").read_text() + "repair: changed\n")
    with pytest.raises(StateConflictError, match="changed since preview/queue"):
        analysis_resume.verify(context, analysis_resume.inspect(attempt))


@pytest.mark.parametrize(
    "damage", ["missing", "owner", "samples", "project", "link", "pep", "libprep", "workflow"]
)
def test_invalid_context_never_mutates_state_or_deletes_outputs(prepared, damage):
    cfg, store, output, work = prepared
    if damage == "missing":
        (work / "config.yaml").unlink()
    elif damage == "owner":
        (work / analysis_snapshots.MARKER).write_text(json.dumps({"run_id": "other"}))
    elif damage in {"samples", "project", "workflow"}:
        config = yaml.safe_load((work / "config.yaml").read_text())
        key = {"samples": "samples", "project": "project_id", "workflow": "workflow"}[damage]
        config[key] = "wrong" if damage != "workflow" else "../bad"
        (work / "config.yaml").write_text(yaml.safe_dump(config))
    elif damage == "link":
        link = next((work / "data/raw/fastq").iterdir())
        link.unlink()
        link.symlink_to(output / "missing.fastq.gz")
    elif damage == "pep":
        (work / "pep/sample_table.csv").unlink()
    else:
        (work / "src/gcf-workflows/libprep.config").unlink()
    before = store.state_path(RUN_ID).read_bytes()
    with pytest.raises(StateConflictError, match="Cannot resume"):
        queue()
    assert store.state_path(RUN_ID).read_bytes() == before
    assert (work / ".snakemake/metadata").read_text() == "provenance\n"


@pytest.mark.parametrize(
    "options",
    [{"from_stage": "reporting"}, {"from_stage": "demultiplexing"}, {"refresh_inputs": True}],
)
def test_incompatible_options_rejected_before_state_access(options, monkeypatch):
    get_cfg = Mock(side_effect=AssertionError("configuration must not be loaded"))
    monkeypatch.setattr(manager, "get_cfg", get_cfg)
    with pytest.raises(StateConflictError, match="--resume requires"):
        manager.rerun_flowcell(
            flowcell=RUN_ID, resume=True, **{"from_stage": "analysis", **options}
        )
    get_cfg.assert_not_called()


def test_active_lease_not_bypassed_by_force(prepared):
    cfg, store, output, work = prepared
    with store.execution_lease(RUN_ID), pytest.raises((ExecutionLeaseError, StateConflictError)):
        queue()


def test_resume_skips_regeneration_and_refreshes_delivery_copies(prepared, monkeypatch):
    cfg, store, output, work = prepared
    context = analysis_resume.inspect(store.read(RUN_ID))
    cfg.run.pipeline = "singlecell"
    before = bytes_and_times(work)
    subprocess = Mock(side_effect=AssertionError("configmaker must not run"))
    monkeypatch.setattr(afterFastq.subprocess, "check_call", subprocess)
    select = Mock(side_effect=AssertionError("installed libprep must not be selected"))
    monkeypatch.setattr(afterFastq, "select_workflow", select)
    run = Mock()
    monkeypatch.setattr(afterFastq, "run_logged_command", run)
    stale = output / f"QC_{PROJECT}/bfq/stale.txt"
    stale.parent.mkdir(parents=True)
    stale.write_text("old")
    summary = output / f"all_samples_web_summary_{PROJECT}_260918.html"
    summary.write_text("old")
    afterFastq.full_align(cfg, resume=context)
    assert bytes_and_times(work) == before
    assert not stale.exists() and not summary.exists()
    assert (output / f"QC_{PROJECT}/bfq/multiqc_{PROJECT}.html").read_text() == "current report"
    command = run.call_args.args[0]
    assert "--rerun-incomplete" in command and "--keep-incomplete" not in command
    assert not any(flag.startswith("--force") or flag == "--rerun-triggers" for flag in command)
    assert run.call_args.kwargs["cwd"] == work


def test_valid_multi_project_resume_and_missing_second_project(prepared):
    cfg, store, output, work = prepared
    other = "GCF-2026-002"
    sheet = output / "SampleSheet.csv"
    sheet.write_text(sheet.read_text() + f"sample2,{other},TGCA\n")
    book = load_workbook(output / "Sample-Submission-Form.xlsx")
    book["Sample-Submission-Form"].cell(17, 1, "sample2")
    book["INFO (GCF-lab only)"].append(["sample2"])
    book.save(output / "Sample-Submission-Form.xlsx")
    fastq = output / other / "sample2_R1.fastq.gz"
    write_fastq(fastq)
    work2 = analysis_snapshots.workdir_path(RUN_ID, other)
    shutil.copytree(work, work2, symlinks=True)
    identity = analysis_snapshots.identify_workdir(work2, RUN_ID, other)
    config = yaml.safe_load((work2 / "config.yaml").read_text())
    config["project_id"] = [other]
    config["samples"] = {"sample2": dict(config["samples"]["sample"], Sample_ID="sample2")}
    (work2 / "config.yaml").write_text(yaml.safe_dump(config))
    (work2 / "pep/sample_table.csv").write_text("sample_name\nsample2\n")
    (work2 / "data/raw/fastq/sample_R1.fastq.gz").unlink()
    (work2 / "data/raw/fastq/sample2_R1.fastq.gz").symlink_to(fastq)
    state = store.read(RUN_ID)
    state["projects"].append(other)
    state["stages"]["analysis"]["metadata"]["workdirs"][other] = identity
    store.write(state)
    queued = queue()
    assert set(queued["restart_request"]["analysis_resume"]["projects"]) == {PROJECT, other}
    (work2 / "config.yaml").unlink()
    before = store.state_path(RUN_ID).read_bytes(), bytes_and_times(work)
    with pytest.raises(StateConflictError, match="cannot read"):
        queue()
    assert (store.state_path(RUN_ID).read_bytes(), bytes_and_times(work)) == before


def test_well_repairs_must_match_curated_mapping(prepared):
    cfg, store, output, work = prepared
    selection, result = preflight.require_valid_inputs(store.read(RUN_ID)["source_path"], output)
    result.forms[0]["demux"] = {"data": pd.DataFrame([{"Sample_ID": "LE01", "Wells": "A1-A2"}])}
    config = yaml.safe_load((work / "config.yaml").read_text())
    config["wells"] = {"ÅLE01": {"Sample_ID": "ÅLE01", "Wells": "A1-A2"}}
    (work / "config.yaml").write_text(yaml.safe_dump(config))
    with pytest.raises(StateConflictError, match="wells IDs differ"):
        analysis_resume.inspect(store.read(RUN_ID), selection=selection, result=result)
    config["wells"] = {"LE01": {"Sample_ID": "LE01", "Wells": "A1-A2"}}
    (work / "config.yaml").write_text(yaml.safe_dump(config))
    context = analysis_resume.inspect(store.read(RUN_ID), selection=selection, result=result)
    assert context["mode"] == "resume"
    config["wells"]["LE01"]["Wells"] = "A3-A4"
    (work / "config.yaml").write_text(yaml.safe_dump(config))
    with pytest.raises(StateConflictError, match="wells mapping differs"):
        analysis_resume.inspect(store.read(RUN_ID), selection=selection, result=result)


def test_lane_repeated_identity_is_valid_but_foreign_flowcell_is_not(prepared):
    cfg, store, output, work = prepared
    config = yaml.safe_load((work / "config.yaml").read_text())
    sample = config["samples"]["sample"]
    sample["Flowcell_ID"] = ",".join([RUN_ID.split("_")[-1]] * 4)
    sample["Flowcell_Name"] = ",".join([RUN_ID] * 4)
    (work / "config.yaml").write_text(yaml.safe_dump(config))
    assert analysis_resume.inspect(store.read(RUN_ID))["mode"] == "resume"
    sample["Flowcell_Name"] += ",another_run"
    (work / "config.yaml").write_text(yaml.safe_dump(config))
    with pytest.raises(StateConflictError, match="inconsistent Flowcell_Name"):
        queue()
