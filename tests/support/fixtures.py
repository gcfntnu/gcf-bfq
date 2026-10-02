"""Small shared fixtures extracted from the existing state integration tests."""

import gzip

from bcl2fastq_pipeline.config import Paths, PipelineConfig, RunContext, StaticConfig
from bcl2fastq_pipeline.state import new_state
from openpyxl import Workbook

RUN_ID = "260918_MN00686_0026_A000HCMFHF"


def configured_bfq(tmp_path, run_id=RUN_ID):
    paths = Paths(
        ekista_base_dir=tmp_path / "ekista",
        nova_base_dir=tmp_path / "nova",
        output_dir=tmp_path / "output",
        log_dir=tmp_path / "logs",
        manager_dir=tmp_path / "manager",
        report_dir=tmp_path / "reports",
        analysis_dir=tmp_path / "analysis",
    )
    for path in (
        paths.ekista_base_dir,
        paths.nova_base_dir,
        paths.output_dir,
        paths.log_dir,
        paths.manager_dir,
        paths.report_dir,
        paths.analysis_dir,
    ):
        path.mkdir(parents=True, exist_ok=True)
    cfg = PipelineConfig(
        static=StaticConfig(
            paths=paths,
            system={"minspace": "1", "sleeptime": "1"},
            version={"pipeline": "0.3.1"},
        ),
        run=RunContext(),
    )
    PipelineConfig._instance = cfg
    source = paths.nova_base_dir / run_id
    source.mkdir()
    cfg.run.begin(source, paths)
    return cfg, source, paths.output_dir / run_id


def write_inputs(directory, suffix=""):
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "SampleSheet.csv").write_text(
        f"[CustomOptions]\nLibprep,Illumina DNA Prep\nUser,test{suffix}\n"
        "[Data]\nSample_ID,Sample_Project,index\nsample,GCF-2026-001,ACGT\n",
        encoding="utf-8",
    )
    workbook = Workbook()
    customer = workbook.active
    customer.title = "Sample-Submission-Form"
    customer.cell(15, 1, "Unique Sample ID")
    customer.cell(15, 2, "Sample Group")
    customer.cell(16, 1, "sample")
    customer.cell(16, 2, f"group{suffix}")
    lab = workbook.create_sheet("INFO (GCF-lab only)")
    lab.append(["Sample_ID"])
    lab.append(["sample"])
    workbook.save(directory / "Sample-Submission-Form.xlsx")


def write_fastq(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt") as handle:
        handle.write("@read\nACGT\n+\n!!!!\n")


def completed_state(cfg, source, output):
    state = new_state(
        RUN_ID,
        source,
        output,
        origin="new",
        start_stage="demultiplexing",
        cfg=cfg,
    )
    state["status"] = "completed"
    state["current_stage"] = "finalization"
    state["completed_at"] = "2026-09-27T12:00:00+00:00"
    state["projects"] = ["GCF-2026-001"]
    for detail in state["stages"].values():
        detail["status"] = "completed"
        detail["completed_at"] = "2026-09-27T12:00:00+00:00"
    return state
